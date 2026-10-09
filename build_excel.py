# -*- coding: utf-8 -*-
"""build_excel.py — Generate laporan Excel Ado-BranchBekasi dari js/data.js.
Sheet: Ringkasan (KPI + per STO), Detail ODC (676 baris + list pelanggan),
Detail Feeder (244 kabel). Jalankan setelah build_data.py:
  "C:/Users/syahr/AppData/Local/Programs/Python/Python314/python.exe" build_excel.py
"""
import json, os, re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(BASE, 'js', 'data.js'), encoding='utf-8') as f:
    src = f.read()
ADO = json.loads(re.search(r'const ADO = (\{.*\});', src, re.S).group(1))

T = ADO['totals']; M = ADO['meta']

# ---- palet gaya dashboard ----
MAROON = '9E1B32'
RED    = 'C0392B'
TEAL   = '0D6E6A'
SKY    = '2874A6'
CREAM  = 'FAF7F5'
MUTED  = 'F3E9E7'
BORDER_C = 'EAD5D0'
INK    = '3B1219'
WARNBG = 'FFF9EE'

def F(bold=False, size=10, color=INK, name='Fira Sans'):
    return Font(name=name, bold=bold, size=size, color=color)
def fill(hex_):
    return PatternFill('solid', fgColor=hex_)
AL = Alignment(horizontal='left', vertical='center')
AR = Alignment(horizontal='right', vertical='center')
AC = Alignment(horizontal='center', vertical='center')
WRAP = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Side(style='thin', color=BORDER_C)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUMFMT = '#,##0'
PCTFMT = '0.0%'

wb = Workbook()

# ================= SHEET 1 — RINGKASAN =================
ws = wb.active
ws.title = 'Ringkasan'
ws.sheet_view.showGridLines = False

ws.merge_cells('A1:M1')
ws['A1'] = 'Ado-BranchBekasi — Analisa Daerah Operasi Branch Bekasi'
ws['A1'].font = F(bold=True, size=16, color=MAROON)
ws.merge_cells('A2:M2')
ws['A2'] = 'Bidang Bisnis TelkomAkses · Sumber: ALPOR ' + M['dataDate'] + ' · dibuat ' + datetime.now().strftime('%d %b %Y %H:%M')
ws['A2'].font = F(size=10, color='6B5A5E')

# blok KPI
kpis = [
    ('STO', T['n_sto'], 'di 2 area (Kota + Depok)', MAROON),
    ('FEEDER', T['n_feeder'], f"{T['n_seg']:,} segmen".replace(',', '.'), TEAL),
    ('ODC', T['n_odc'], 'unit aktif', RED),
    ('KAPASITAS FEEDER (CORE)', T['fe_cap'], f"used {T['fe_used']:,} · idle {T['fe_idle']:,}".replace(',', '.'), SKY),
    ('KAPASITAS ODC (PORT)', T['odc_kap_akt'], f"used {T['odc_used']:,} · idle {T['odc_idle']:,}".replace(',', '.'), MAROON),
    ('LIST PELANGGAN', T['pelanggan'], 'total branch Bekasi', TEAL),
]
row = 4
for i, (lbl, val, sub, colr) in enumerate(kpis):
    r = row + i
    ws.merge_cells(f'A{r}:C{r}')
    ws[f'A{r}'] = lbl
    ws[f'A{r}'].font = F(bold=True, size=9, color='6B5A5E')
    ws.merge_cells(f'D{r}:F{r}')
    ws[f'D{r}'] = val
    ws[f'D{r}'].font = F(bold=True, size=14, color=colr, name='Fira Code')
    ws[f'D{r}'].number_format = NUMFMT
    ws[f'D{r}'].alignment = AR
    ws[f'A{r}'].fill = fill(MUTED); ws[f'B{r}'].fill = fill(MUTED); ws[f'C{r}'].fill = fill(MUTED)
    ws[f'A{r}'].border = BOX; ws[f'B{r}'].border = BOX; ws[f'C{r}'].border = BOX
    ws[f'D{r}'].border = BOX; ws[f'E{r}'].border = BOX; ws[f'D{r}'].border = BOX

# tabel per STO
TR = 11
headers = ['STO', 'Area', 'Feeder\nJml', 'Feeder\nKap (core)', 'Feeder\nUsed', 'Feeder\nIdle',
           'Okup. FE', 'ODC\nJml', 'ODC\nKap (port)', 'ODC\nUsed', 'ODC\nIdle', 'Okup. ODC', 'List Pelanggan']
for ci, h in enumerate(headers, 1):
    c = ws.cell(row=TR, column=ci, value=h)
    c.font = F(bold=True, size=9, color='FFFFFF')
    c.fill = fill(MAROON)
    c.alignment = AC
    c.border = BOX
for si, s in enumerate(ADO['sto']):
    r = TR + 1 + si
    vals = [s['code'],
            'Bekasi Kota' if s['area'] == 'BEKASI KOTA' else 'Bekasi Depok',
            s['n_feeder'], s['fe_cap'], s['fe_used'], s['fe_idle'],
            (s['fe_occ'] / 100 if s['fe_occ'] is not None else None),
            s['n_odc'], s['odc_kap_akt'], s['odc_used'], s['odc_idle'],
            (s['odc_occ'] / 100 if s['odc_occ'] is not None else None),
            s['pelanggan']]
    for ci, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=ci, value=v)
        c.border = BOX
        c.font = F()
        if ci == 1:
            c.font = F(bold=True, color=s['color'].lstrip('#').upper() if s['color'] != '#C0392B' else RED)
            c.font = Font(name='Fira Code', bold=True, size=10, color=s['color'].lstrip('#'))
            c.alignment = AL
        elif ci == 2:
            c.alignment = AL
        else:
            c.alignment = AR
            if ci in (7, 12):
                c.number_format = PCTFMT
            else:
                c.number_format = NUMFMT
        if si % 2 == 1:
            c.fill = fill('FDFEFF')
# TOTAL
r = TR + 1 + len(ADO['sto'])
tvals = ['TOTAL', f"{T['n_sto']} STO", T['n_feeder'], T['fe_cap'], T['fe_used'], T['fe_idle'],
         T['fe_occ'] / 100, T['n_odc'], T['odc_kap_akt'], T['odc_used'], T['odc_idle'],
         T['odc_occ'] / 100, T['pelanggan']]
for ci, v in enumerate(tvals, 1):
    c = ws.cell(row=r, column=ci, value=v)
    c.font = F(bold=True, size=10)
    c.fill = fill(MUTED)
    c.border = Border(top=Side(style='double', color='D9BCB4'), left=thin, right=thin, bottom=thin)
    c.alignment = AL if ci <= 2 else AR
    if ci in (7, 12):
        c.number_format = PCTFMT
    elif ci > 2:
        c.number_format = NUMFMT
    if ci == 13:
        c.font = Font(name='Fira Code', bold=True, size=10, color=MAROON)
# catatan
r += 2
ws.merge_cells(f'A{r}:M{r}')
ws[f'A{r}'] = ('Catatan: kolom "List Pelanggan" per STO dari file list pelanggan branch; rincian alokasi per ODC '
               '(proporsional okupansi port used) ada di sheet "Detail ODC". Okupansi FE = used/kapasitas core, '
               'Okupansi ODC = port used/kapasitas aktual.')
ws[f'A{r}'].font = F(size=9, color='6B5A5E')
ws[f'A{r}'].alignment = WRAP
ws.row_dimensions[r].height = 28
widths = [7, 13, 9, 12, 11, 11, 9, 9, 12, 11, 11, 9, 14]
for ci, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.freeze_panes = f'A{TR + 1}'

# ================= SHEET 2 — DETAIL ODC =================
ws2 = wb.create_sheet('Detail ODC')
ws2.sheet_view.showGridLines = False
h2 = ['STO', 'Nama ODC', 'Spec', 'Kap. Potensial', 'Kap. Aktual', 'Port Used', 'Port Idle',
      'Okupansi', 'Status', 'List Pelanggan']
for ci, h in enumerate(h2, 1):
    c = ws2.cell(row=1, column=ci, value=h)
    c.font = F(bold=True, size=9, color='FFFFFF')
    c.fill = fill(MAROON); c.alignment = AC; c.border = BOX
codes = {s['code']: s for s in ADO['sto']}
ri = 2
for o in sorted(ADO['odc'], key=lambda x: (list(codes).index(x[0]), -x[9])):
    sto, name, spec, kap_akt, used, lat, lon, flags, kap_pot, pel = o
    colr = codes[sto]['color'].lstrip('#')
    idle = (kap_akt or 0) - (used or 0)
    occ = (used / kap_akt) if (kap_akt and used is not None) else None
    status = []
    if flags & 1: status.append('data kotor')
    if flags & 2: status.append('kap. aktual > potensial')
    if not status: status = ['OK']
    vals = [sto, name, spec, kap_pot, kap_akt, used, idle, occ, ' / '.join(status), pel]
    for ci, v in enumerate(vals, 1):
        c = ws2.cell(row=ri, column=ci, value=v)
        c.border = BOX; c.font = F()
        if ci == 1:
            c.font = Font(name='Fira Code', bold=True, size=9, color=colr); c.alignment = AL
        elif ci == 2:
            c.font = Font(name='Fira Code', size=9); c.alignment = AL
        elif ci == 3:
            c.alignment = AL
        elif ci == 8:
            c.alignment = AR; c.number_format = PCTFMT
        elif ci == 9:
            c.alignment = AL
            if v == 'OK':
                c.font = F(size=9, color='166534')
        elif ci == 10:
            c.font = Font(name='Fira Code', size=9, color=MAROON); c.alignment = AR; c.number_format = NUMFMT
        else:
            c.alignment = AR; c.number_format = NUMFMT
        if ri % 2 == 0:
            c.fill = fill('FDFEFF')
    ri += 1
# total
for ci, v in enumerate(['TOTAL', f'{ri-2} ODC', '', None, None, None, None, None, '', T['pelanggan']], 1):
    c = ws2.cell(row=ri, column=ci, value=v)
    c.font = F(bold=True, size=9); c.fill = fill(MUTED); c.alignment = AR if ci > 1 else AL
    c.border = Border(top=Side(style='double', color='D9BCB4'), left=thin, right=thin, bottom=thin)
    if ci == 10:
        c.font = Font(name='Fira Code', bold=True, size=9, color=MAROON); c.number_format = NUMFMT
w2 = [7, 16, 13, 13, 12, 11, 11, 10, 22, 14]
for ci, w in enumerate(w2, 1):
    ws2.column_dimensions[get_column_letter(ci)].width = w
ws2.freeze_panes = 'C2'
ws2.auto_filter.ref = f'A1:J{ri-1}'

# ================= SHEET 3 — DETAIL FEEDER =================
ws3 = wb.create_sheet('Detail Feeder')
ws3.sheet_view.showGridLines = False
h3 = ['STO', 'Kabel Feeder', 'FTM', 'Spec', 'Kapasitas (core)', 'Used', 'Idle', 'Spare', 'Okupansi', 'Status']
for ci, h in enumerate(h3, 1):
    c = ws3.cell(row=1, column=ci, value=h)
    c.font = F(bold=True, size=9, color='FFFFFF')
    c.fill = fill(TEAL); c.alignment = AC; c.border = BOX
ri = 2
for fe in ADO['feeders']:
    sto, name, spec, cap, used, idle, spare, flags, segs, ftm = fe
    colr = codes[sto]['color'].lstrip('#')
    occ = (used / cap) if (cap and used is not None) else None
    status = []
    if flags & 1: status.append('no FTM/kapasitas')
    if flags & 2: status.append('used+idle+spare ≠ kapasitas')
    if not status: status = ['OK']
    vals = [sto, name, ftm, spec, cap, used, idle, spare, occ, ' / '.join(status)]
    for ci, v in enumerate(vals, 1):
        c = ws3.cell(row=ri, column=ci, value=v)
        c.border = BOX; c.font = F()
        if ci == 1:
            c.font = Font(name='Fira Code', bold=True, size=9, color=colr); c.alignment = AL
        elif ci == 2:
            c.font = Font(name='Fira Code', size=9); c.alignment = AL
        elif ci in (3, 4):
            c.alignment = AL
        elif ci == 9:
            c.alignment = AR; c.number_format = PCTFMT
        elif ci == 10:
            c.alignment = AL
            if v == 'OK':
                c.font = F(size=9, color='166534')
        else:
            c.alignment = AR; c.number_format = NUMFMT
        if ri % 2 == 0:
            c.fill = fill('FDFEFF')
    ri += 1
for ci, v in enumerate(['TOTAL', f'{ri-2} kabel', '', None,
                        sum(f[3] or 0 for f in ADO['feeders']),
                        sum(f[4] or 0 for f in ADO['feeders']),
                        sum(f[5] or 0 for f in ADO['feeders']),
                        sum(f[6] or 0 for f in ADO['feeders']),
                        T['fe_occ'] / 100, ''], 1):
    c = ws3.cell(row=ri, column=ci, value=v)
    c.font = F(bold=True, size=9); c.fill = fill(MUTED); c.alignment = AR if ci > 1 else AL
    c.border = Border(top=Side(style='double', color='D9BCB4'), left=thin, right=thin, bottom=thin)
    if ci == 9:
        c.number_format = PCTFMT
    elif ci > 4:
        c.number_format = NUMFMT
w3 = [7, 15, 14, 13, 14, 10, 10, 10, 10, 26]
for ci, w in enumerate(w3, 1):
    ws3.column_dimensions[get_column_letter(ci)].width = w
ws3.freeze_panes = 'C2'
ws3.auto_filter.ref = f'A1:J{ri-1}'

out = r'C:/Users/syahr/Documents/PROJECT ADO/Ado-BranchBekasi.xlsx'

# page setup: landscape + fit 1 halaman lebar (untuk print/PDF)
for sheet in [ws, ws2, ws3]:
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True

wb.save(out)
print('SAVED', out)
print('sheets:', wb.sheetnames)
